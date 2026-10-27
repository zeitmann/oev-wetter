"""Folien fuer Paket B (Ist-Daten, Ingestion, Storage, Datenmodell) erstellen.

Aufbau, Farben und Schrift wie bei den Folien von Paket A und C
(src/build_folien_paket_a.py, src/build_folien.py), damit die drei Teile
in der Gesamtpraesentation (src/build_praesentation.py) einheitlich wirken.
Inhalte: docs/quellenprofil_istdaten.md, docs/storage_evaluation.md und
docs/schema_absprache.md; Ankunftsstatus aus docs/ergebnisse/ankunftsstatus.csv.

Optional: pip install python-pptx
Quellen (APA 7):
Canny, S. (2025). python-pptx documentation. https://python-pptx.readthedocs.io/
SBB AG, Geschaeftsstelle Systemaufgaben Kundeninformation. (2026). Ist-Daten v2
    [Datensatz]. Open-Data-Plattform Mobilitaet Schweiz.
    https://archive.opentransportdata.swiss/actual_data_archive.htm
DuckDB Foundation. (2026). DuckDB documentation. https://duckdb.org/docs/
"""
import unicodedata
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
NAVY = '17364A'
GREEN = '167D8D'
GREY = '53616B'
WHITE = 'FFFFFF'
LIGHT = 'EEF4F6'
FOOTER = ('Gruppe 4 · Aladin · Paket B | Quellen: SBB/SKI (2026), '
          'swisstopo (2026), DuckDB Foundation (2026)')


def out_dir():
    """Ordner docs/praesentation (Umlaut-Schreibweise des Dateisystems egal)."""
    for p in (ROOT / 'docs').iterdir():
        if unicodedata.normalize('NFC', p.name) == 'präsentation':
            return p
    p = ROOT / 'docs' / 'präsentation'
    p.mkdir(parents=True, exist_ok=True)
    return p


def text(slide, x, y, w, h, value, size=20, color=NAVY, bold=False, space=12, font='Aptos'):
    """Textfeld; jede Zeile von value wird ein eigener Absatz."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(space)
    return box


def card(slide, x, y, w, h, fill=WHITE, line=GREEN):
    """Abgerundete Karte wie bei Paket A."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def arrow(slide, x1, y1, x2, y2):
    """Verbindungslinie mit Pfeilspitze (fuer Ablauf und Datenmodell)."""
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = RGBColor.from_string(GREEN)
    line.line.width = Pt(2)
    # Pfeilspitze am Ende: python-pptx hat dafuer keine API, darum direkt im XML
    ln = line.line._get_or_add_ln()
    tail = ln.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'})
    ln.append(tail)
    return line


def table(slide, x, y, w, rows, widths, size=15, highlight=None):
    """Einfache Tabelle: erste Zeile Kopf (Navy), optional eine Zeile hervorgehoben."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(.45 * len(rows)))
    tbl = shape.table
    for j, width in enumerate(widths):
        tbl.columns[j].width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = value
            cell.fill.solid()
            if i == 0:
                fill, color, bold = NAVY, WHITE, True
            elif i == highlight:
                fill, color, bold = GREEN, WHITE, True
            else:
                fill, color, bold = (LIGHT if i % 2 == 0 else WHITE), NAVY, False
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            p = cell.text_frame.paragraphs[0]
            p.font.name = 'Aptos'
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
    return shape


def base(prs, title, index, notes):
    """Leere Folie mit Balken, Titel, Fusszeile, Seitenzahl und Sprechernotizen."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(GREEN)
    bar.line.fill.background()
    text(slide, .5, .35, 12.3, .6, title, 30, bold=True)
    text(slide, .5, 7.08, 11.9, .3, FOOTER, 10, GREY)
    text(slide, 12.4, 7.06, .5, .3, str(index), 11, GREY)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = 'ÖV und Wetter: Paket B'
    prs.core_properties.author = 'Aladin · Gruppe 4'
    prs.core_properties.subject = 'Ist-Daten, Ingestion, Storage und Datenmodell'

    # 1 Datenquelle
    s = base(prs, 'Ist-Daten v2: jeder Halt mit Soll- und Ist-Zeit', 1,
        'Ca. 45 Sekunden. Die Ist-Daten enthalten für jeden Halt jeder Fahrt die Fahrplanzeit und '
        'die effektive Zeit oder die letzte Prognose. Version 2 gibt es ab Juli 2025, unser '
        'Zeitraum ist also vollständig abgedeckt. Die Zeiten sind Lokalzeit ohne Zeitzone, und die '
        'Sommerzeitumstellung am 29. März liegt im Zeitraum. Das behandeln wir beim Laden. '
        'Details: docs/quellenprofil_istdaten.md.')
    table(s, .55, 1.35, 12.2, [
        ['Merkmal', 'Befund'],
        ['Herkunft', 'Geschäftsstelle SKI (SBB AG) im Auftrag des BAV, opentransportdata.swiss'],
        ['Zugang und Lizenz', 'Monatsarchiv als Zip, frei ohne Login, Quellenangabe Pflicht'],
        ['Abdeckung', 'v2 ab Juli 2025, 1.1. bis 31.3.2026 vollständig (90 Tage)'],
        ['Format', 'CSV mit Trennzeichen «;», UTF-8, 22 Spalten, eine Datei pro Tag'],
        ['Grösse', '420 bis 650 MB pro Tag, ca. 2,4 Mio. Zeilen (ganze Schweiz)'],
        ['Zeitformat', 'Lokalzeit ohne Zeitzone, Sommerzeit am 29.3.2026 im Zeitraum'],
    ], [3.0, 9.2], size=17)
    text(s, .55, 5.2, 12.2, .9,
        'Wichtige Spalten: BETRIEBSTAG, PRODUKT_ID, LINIEN_TEXT, BPUIC, ANKUNFTSZEIT, AN_PROGNOSE, '
        'AN_PROGNOSE_STATUS, FAELLT_AUS_TF, DURCHFAHRT_TF', 15, GREY, font='Consolas')

    # 2 Ingestion
    s = base(prs, 'Ingestion: ein Tag nach dem anderen, wiederholbar', 2,
        'Ca. 60 Sekunden. Das Skript lädt Tag für Tag. Die Datei wird nie ganz in den Speicher '
        'geladen: DuckDB liest sie streamend und filtert direkt per SQL. Die gefilterten Rohdaten '
        'legen wir unverändert ab, dann wird der Tag in einer Transaktion in der Datenbank ersetzt. '
        'Dadurch ist das Skript idempotent: Ein zweiter Lauf überspringt fertige Tage, getestet mit '
        'allen 90 Tagen. Downloads landen zuerst in einer Teildatei, Fehler stehen im Log und das '
        'Skript endet mit Fehlercode.')
    schritte = [('1 Laden', 'Monats-Zip aus dem Archiv, mit Wiederholung'),
                ('2 Lesen', 'Tagesdatei streamend mit DuckDB'),
                ('3 Filtern', 'Bahn und Bus, 2587 Haltestellen'),
                ('4 Ablegen', 'gefilterte Rohdaten als CSV'),
                ('5 Speichern', 'Tag in DuckDB ersetzen, eine Transaktion')]
    for i, (head, body) in enumerate(schritte):
        x = .55 + i * 2.5
        last = i == len(schritte) - 1
        card(s, x, 1.4, 2.1, 1.9, fill=NAVY if last else WHITE, line=None if last else GREEN)
        text(s, x + .12, 1.55, 1.9, .5, head, 18, WHITE if last else GREEN, True)
        text(s, x + .12, 2.1, 1.9, 1.1, body, 14, WHITE if last else NAVY, space=4)
        if not last:
            arrow(s, x + 2.12, 2.35, x + 2.48, 2.35)
    merkmale = [('Idempotent', 'Geladene Tage werden übersprungen, ein neuer Lauf erzeugt keine Duplikate.'),
                ('Protokolliert', 'Pro Tag Zeilen vor und nach dem Filter, Dauer und Fehler in Logfile und Tabelle ingest_log.'),
                ('Zeitzonen-sicher', 'Alle Zeiten als TIMESTAMPTZ (ISO 8601 mit Zeitzone), Sommerzeit korrekt.')]
    for i, (head, body) in enumerate(merkmale):
        x = .55 + i * 4.15
        text(s, x, 3.75, 3.9, .5, head, 20, GREEN, True)
        text(s, x, 4.3, 3.9, 1.4, body, 16)
    text(s, .55, 6.1, 12.2, .5,
        'python src/ingest_istdaten.py --start 2026-01-01 --end 2026-03-31 '
        '--stops data/stammdaten/haltestellen_ostschweiz.csv', 13, GREY, font='Consolas')

    # 3 Filter und Datenqualitaet
    s = base(prs, 'Von 215 Mio. auf 17 Mio. Zeilen', 3,
        'Ca. 60 Sekunden. Ungefiltert wären es 215 Millionen Zeilen. Wir behalten nur Bahn und Bus '
        'an Haltestellen in den vier Kantonen. Die Haltestellenliste haben wir selbst gebaut: '
        'Koordinaten aus dem GTFS-Fahrplan, Kanton per Punkt-in-Polygon mit den Kantonsgrenzen von '
        'swisstopo. Übrig bleiben 17 Millionen Halte, fast 90 Prozent davon Bus. Beim Laden '
        'bereinigen wir nur Technisches (Duplikate, Haltestellennummer, Zeitzone). Inhaltliche '
        'Filter macht Paket C, damit nachvollziehbar bleibt, was wegfällt. Wichtig für die Analyse: '
        'Bei der Bahn sind knapp 90 Prozent der Ankünfte echt gemessen, beim Bus nur zwei Drittel.')
    zahlen = [('215 Mio.', 'Zeilen roh, 90 Tage, ganze Schweiz', False),
              ('2587', 'Haltestellen in SG, TG, AR und AI', False),
              ('17,07 Mio.', 'Halte in der Datenbank (Bahn 1,75, Bus 15,32)', True)]
    for i, (zahl, body, dark) in enumerate(zahlen):
        x = .55 + i * 4.15
        card(s, x, 1.35, 3.95, 2.0, fill=NAVY if dark else LIGHT, line=None)
        text(s, x + .25, 1.5, 3.5, .9, zahl, 40, WHITE if dark else NAVY, True)
        text(s, x + .25, 2.45, 3.5, .8, body, 16, WHITE if dark else GREY)
    text(s, .55, 3.6, 6.0, .5, 'Beim Laden bereinigt', 20, GREEN, True)
    text(s, .55, 4.15, 6.0, 2.2,
        '• Exakte Duplikate entfernt, gemessener Wert vor Prognose\n'
        '• 9-stellige BPUIC auf die 7-stellige Haltestelle gekürzt\n'
        '• Lokalzeit in TIMESTAMPTZ (Europe/Zurich) umgewandelt', 16, space=6)
    text(s, 6.85, 3.6, 6.0, .5, 'Offengelegt für Paket C', 20, GREEN, True)
    text(s, 6.85, 4.15, 6.0, 2.2,
        '• Ankunft echt gemessen (REAL): Bahn 89,7 %, Bus 65,9 %\n'
        '• Unplausible Ist-Zeiten (z. B. genau 1 Tag daneben) bleiben roh\n'
        '• Status jeder Zeit (REAL, PROGNOSE, GESCHAETZT, UNBEKANNT, leer) wird mitgeführt', 16, space=6)

    # 4 Storage-Evaluation
    s = base(prs, 'Storage: DuckDB – analytisch, eine Datei, kein Server', 4,
        'Ca. 60 Sekunden. Unsere Fragen sind analytisch: gruppieren und verknüpfen über Millionen '
        'Zeilen. Dafür ist eine spaltenorientierte Datenbank wie DuckDB gebaut. Sie liest die '
        'grossen CSV-Dateien direkt und streamend und ist eine einzige Datei ohne Server, also kann '
        'jedes Gruppenmitglied sie sofort nachbauen. Die Schwäche ist der Mehrbenutzerbetrieb. '
        'Falls das Dashboard in Teil 2 viele gleichzeitige Zugriffe braucht, wechseln wir auf '
        'PostgreSQL; das Schema ist Standard-SQL. Details: docs/storage_evaluation.md.')
    table(s, .55, 1.35, 12.2, [
        ['Kriterium', 'DuckDB', 'SQLite', 'PostgreSQL', 'MongoDB'],
        ['Analytische Abfragen (Joins, GROUP BY)', 'sehr gut', 'mässig', 'gut', 'mässig'],
        ['CSV streamend einlesen', 'sehr gut', 'mässig', 'gut', 'mässig'],
        ['Zeitzonen (TIMESTAMPTZ)', 'gut', 'schwach', 'sehr gut', 'mässig'],
        ['Skalierung auf mehr Kantone', 'gut', 'schwach', 'sehr gut', 'gut'],
        ['Setup für die Gruppe', 'sehr gering', 'sehr gering', 'hoch (Server)', 'hoch (Server)'],
        ['Viele gleichzeitige Nutzer', 'mässig', 'mässig', 'sehr gut', 'gut'],
    ], [4.2, 2.0, 2.0, 2.0, 2.0], size=16)
    text(s, .55, 5.0, 12.2, 1.2,
        'Entscheid: DuckDB für Teil 1. Für das Dashboard in Teil 2 ist ein Wechsel auf PostgreSQL '
        'möglich, weil das Schema Standard-SQL ist.', 19, GREEN, True)

    # 5 Datenmodell
    s = base(prs, 'Datenmodell: Halt und Wetter über die Haltestelle verknüpft', 5,
        'Ca. 50 Sekunden. Die Tabelle halt ist das Herz: ein Eintrag pro Halt einer Fahrt, mit '
        'Soll, Ist, Status und Verspätung. Über die 7-stellige Haltestellennummer hängt sie an der '
        'Haltestelle. Paket C ordnet jeder Haltestelle die nächste Wetterstation zu, und über '
        'Station und Stunde kommen die Messwerte dazu. Das ingest_log protokolliert jeden '
        'Ladelauf. Schema und Spaltennamen haben B und C am Anfang gemeinsam festgelegt '
        '(docs/schema_absprache.md); Indizes auf Soll-Zeit und Haltestelle.')
    boxen = [('halt', 'Halt einer Fahrt: Soll, Ist, Status, Verspätung · 17,07 Mio.', NAVY),
             ('haltestelle', '2587 Haltestellen mit Kanton und Koordinaten', WHITE),
             ('haltestelle_station', 'nächste Wetterstation pro Haltestelle (C)', WHITE),
             ('wetterstation', 'MeteoSchweiz-Stationen mit Höhe (C)', WHITE)]
    for i, (name, body, fill) in enumerate(boxen):
        x = .55 + i * 3.15
        dark = fill == NAVY
        card(s, x, 1.5, 2.75, 1.9, fill=fill, line=None if dark else GREEN)
        text(s, x + .15, 1.65, 2.5, .5, name, 14, WHITE if dark else GREEN, True, font='Consolas')
        text(s, x + .15, 2.2, 2.5, 1.1, body, 14, WHITE if dark else NAVY, space=4)
        if i < len(boxen) - 1:
            arrow(s, x + 2.77, 2.45, x + 3.13, 2.45)
    # halt.bpuic ist meist 9-stellig (Kante); gejoint wird ueber die 7-stellige bpuic_haltestelle
    text(s, 1.9, 3.45, 3.0, .4, 'bpuic_haltestelle = bpuic', 12, GREY, font='Consolas')
    text(s, 6.1, 3.45, 3.0, .4, 'bpuic', 12, GREY, font='Consolas')
    text(s, 9.25, 3.45, 3.0, .4, 'station_abk', 12, GREY, font='Consolas')
    unten = [(.55, 'ingest_log', 'Pro Tag: Status, Zeilen vor/nach Filter, Dauer, Fehler'),
             (9.95, 'wettermessung', 'Messwert pro Station und Stunde (UTC) (C)')]
    for x, name, body in unten:
        card(s, x, 4.2, 2.75, 1.7)
        text(s, x + .15, 4.35, 2.5, .5, name, 16, GREEN, True, font='Consolas')
        text(s, x + .15, 4.9, 2.5, 1.0, body, 14, space=4)
    arrow(s, 11.32, 3.4, 11.32, 4.2)
    text(s, 3.6, 4.5, 6.1, 1.4,
        'Alle Zeiten TIMESTAMPTZ · Schlüssel und Spaltennamen mit Paket C abgestimmt · '
        'Indizes auf Soll-Zeit und Haltestelle', 16, GREY)

    # 6 Skalierbarkeit
    s = base(prs, 'Skalierbarkeit: gleicher Code, längere Haltestellenliste', 6,
        'Ca. 45 Sekunden. Die ersten zwei Zeilen haben wir gemessen, die letzte ist hochgerechnet. '
        'Für mehr Kantone braucht es nur eine längere Haltestellenliste, Code und Schema bleiben '
        'gleich. Bei der ganzen Schweiz könnte man die Daten als Parquet liegen lassen und DuckDB '
        'direkt darauf abfragen. Die echte Grenze ist der Mehrbenutzerbetrieb, dafür wäre '
        'PostgreSQL die nächste Stufe. Die gefilterten Daten liegen als Parquet (305 MB) auf '
        'OneDrive, Link im README. Damit übergebe ich an Enis.')
    table(s, .55, 1.35, 12.2, [
        ['Umfang', 'Zeilen', 'DuckDB', 'Parquet', 'Laufzeit'],
        ['1 Tag, ganze Schweiz (gemessen)', '2,34 Mio.', '375 MB', 'ca. 45 MB', '43 s'],
        ['90 Tage, Ostschweiz (gemessen)', '17,07 Mio.', '2,6 GB', '305 MB', 'ca. 16 Min.'],
        ['90 Tage, ganze Schweiz (hochgerechnet)', 'ca. 210 Mio.', 'ca. 34 GB', 'ca. 4 GB', 'ca. 1 h'],
    ], [4.6, 1.9, 1.9, 1.9, 1.9], size=17, highlight=2)
    text(s, .55, 3.7, 12.2, 2.4,
        '• Speicherbedarf beim Laden hängt nicht von der Dateigrösse ab, weil streamend gelesen wird.\n'
        '• Mehr Kantone: nur die Haltestellenliste erweitern, Code und Schema bleiben gleich.\n'
        '• Grenze: DuckDB erlaubt nur einen schreibenden Prozess; für Teil 2 ist PostgreSQL die nächste Stufe.',
        18, space=10)

    target = out_dir() / 'folien_paket_b.pptx'
    prs.save(target)
    print(f'{target.name}: {len(prs.slides)} Folien mit Sprechernotizen')


if __name__ == '__main__':
    main()
