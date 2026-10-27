"""Gesamtpraesentation fuer das Kolloquium aus den Folien von Paket A, B und C.

Liest die drei Foliensaetze in docs/praesentation/ und fuegt sie in der
Reihenfolge des Referats zusammen (A ca. 5 Min., B ca. 5 Min., C ca. 5 Min.,
danach Ausblick und Anhang). Folien, Bilder, Links und Sprechernotizen werden
unveraendert uebernommen. Neu gesetzt werden die Seitenzahlen und das Design:
Hintergrundbild zum Thema (Wolken, Schnee, Zug und Bus, src/build_design.py)
statt des schlichten weissen Hintergrunds mit Balken.
Die Einzelfolien zuerst erzeugen (src/build_folien_paket_a.py,
src/build_folien_paket_b.py, src/build_folien.py, src/build_design.py),
dann dieses Skript:

    python src/build_praesentation.py

Optional: pip install python-pptx
Quellen (APA 7):
Canny, S. (2025). python-pptx documentation. https://python-pptx.readthedocs.io/
"""
import copy
import io

from pptx import Presentation
from pptx.util import Inches

from build_folien_paket_b import out_dir

# Reihenfolge des Referats: (Foliensatz, Foliennummer im Satz, 1-basiert)
ABLAUF = (
    [('a', n) for n in range(1, 7)]      # Titel, Ablauf, Problem, Hypothesen, Methodik, Lifecycle
    + [('b', n) for n in range(1, 7)]    # Quelle, Ingestion, Filter, Storage, Datenmodell, Skalierung
    + [('c', n) for n in range(1, 6)]    # Wetter, Bereinigung, H1/H2, H3, Limitationen
    + [('a', 7)]                         # Ausblick Teil 2, Überleitung zur Diskussion
    + [('a', 8), ('c', 6)]               # Anhang: Quellen
)
DATEIEN = {'a': 'folien_paket_a.pptx', 'b': 'folien_paket_b.pptx', 'c': 'folien_paket_c.pptx'}
RID_ATTRS = ('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',
             '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id',
             '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link')


def copy_slide(src, prs):
    """Eine Folie aus einer anderen Praesentation an prs anhaengen.

    python-pptx kann Folien nicht direkt zwischen Dateien kopieren. Darum werden
    Hintergrund und Formen als XML kopiert und die Beziehungen (Bilder, Links)
    neu angelegt; die alten rId-Verweise werden auf die neuen umgeschrieben.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in list(slide.shapes):          # Layout 6 ist leer, sicherheitshalber
        shape._element.getparent().remove(shape._element)

    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        if rel.is_external:
            rid_map[rid] = slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype.endswith('/image'):
            # Bild als neuen Teil anlegen, sonst kollidieren die Dateinamen beim Speichern
            _, rid_map[rid] = slide.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        else:
            raise ValueError(f'Beziehungstyp nicht unterstuetzt: {rel.reltype}')

    def remap(element):
        for el in element.iter():
            for attr in RID_ATTRS:
                if el.get(attr) in rid_map:
                    el.set(attr, rid_map[el.get(attr)])
        return element

    bg = src._element.cSld.bg
    if bg is not None:
        slide._element.cSld.insert(0, remap(copy.deepcopy(bg)))
    tree = slide.shapes._spTree
    for el in src.shapes._spTree:
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'):
            continue
        tree.append(remap(copy.deepcopy(el)))

    if src.has_notes_slide:
        slide.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return slide


def set_page_number(slide, number):
    """Seitenzahl unten rechts (Textfeld bei x = 12.4 Zoll) neu setzen."""
    for shape in slide.shapes:
        if (shape.has_text_frame and abs(shape.left - Inches(12.4)) < Inches(.05)
                and shape.text_frame.text.strip().isdigit()):
            shape.text_frame.paragraphs[0].runs[0].text = str(number)


def apply_design(slide, image, prs, titel=False):
    """Hintergrundbild hinter alle Formen legen und den gruenen Balken oben entfernen.

    Auf der Titelfolie entfallen zusaetzlich Fusszeile und Seitenzahl, weil dort
    die Illustration (Strasse, Bus) den unteren Rand fuellt.
    """
    for shape in list(slide.shapes):
        if titel and shape.top >= Inches(7.0):
            shape._element.getparent().remove(shape._element)
            continue
        # Balken der Einzelfolien: Rechteck oben links ueber die ganze Breite, 0.12 Zoll hoch
        if (shape.left == 0 and shape.top == 0 and shape.width == prs.slide_width
                and shape.height == Inches(.12)):
            shape._element.getparent().remove(shape._element)
    picture = slide.shapes.add_picture(str(image), 0, 0, prs.slide_width, prs.slide_height)
    tree = slide.shapes._spTree
    tree.remove(picture._element)
    tree.insert(2, picture._element)   # nach nvGrpSpPr und grpSpPr = ganz hinten


def main():
    folder = out_dir()
    quellen = {key: Presentation(str(folder / name)) for key, name in DATEIEN.items()}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = 'Beeinflusst das Wetter die Pünktlichkeit des öV? (Gruppe 4)'
    prs.core_properties.author = 'Timon, Aladin, Enis · Gruppe 4'
    prs.core_properties.subject = 'Projektarbeit Teil 1, Data Engineering HS26'

    design = folder / 'design'
    for number, (key, index) in enumerate(ABLAUF, start=1):
        slide = copy_slide(quellen[key].slides[index - 1], prs)
        set_page_number(slide, number)
        image = 'hintergrund_titel.png' if number == 1 else 'hintergrund_inhalt.png'
        apply_design(slide, design / image, prs, titel=number == 1)

    target = folder / 'praesentation_gruppe4_teil1.pptx'
    prs.save(target)
    print(f'{target.name}: {len(prs.slides)} Folien')


if __name__ == '__main__':
    main()
